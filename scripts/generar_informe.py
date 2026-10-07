#!/usr/bin/env python3
"""Genera el informe único a partir de evidencias reales, sin nuevas peticiones."""
import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
TITLE='Informe de Prueba Sustantiva: Verificación de Controles de Acceso en Endpoints del Sistema – OWASP A01:2025'
TZ=ZoneInfo('America/La_Paz')
datasets=[]
for f in (ROOT/'evidencias').glob('python_*/resultados.json'):
    value=json.loads(f.read_text())
    if len(value['records'])==17 and all('status' in r for r in value['records']):datasets.append((f,value))
py_path,py=max(datasets,key=lambda pair:pair[1]['completed_at'])
records=py['records']
captures=json.loads((ROOT/'capturas/indice_capturas.json').read_text())
assert len(captures)==51
ins={r['id']:json.loads((ROOT/f'evidencias/insomnia/{r["id"]}.response.json').read_text()) for r in records}
hashes=set()
for r in records:
    body=(ROOT/f'evidencias/insomnia/{r["id"]}.body').read_bytes()
    h=hashlib.sha256(body).hexdigest()
    assert h==r['sha256'] and len(body)==r['body_bytes']==7541
    assert r['status']==ins[r['id']]['statusCode']==200
    assert len(ins[r['id']]['requestTestResults'])==4
    assert all(t['status']=='passed' for t in ins[r['id']]['requestTestResults'])
    hashes.add(h)
assert len(hashes)==1
body_hash=next(iter(hashes))

def local(ms):return datetime.fromtimestamp(ms/1000,TZ).strftime('%d/%m/%Y %H:%M:%S %z')
ins_start=local(min(r['created'] for r in ins.values()))
ins_end=local(max(r['created'] for r in ins.values()))
py_start=records[0]['started_at'];py_end=py['completed_at']
py_rel=str(py_path.parent.relative_to(ROOT))

out=ROOT/'informe';out.mkdir(exist_ok=True)
summary={'title':TITLE,'target':py['base'],'phase':'Acceso anónimo exclusivamente',
         'verdict':'No concluyente para OWASP A01:2025',
         'observed':'Las 15 rutas y 2 controles devuelven HTTP 200 text/html, 7541 bytes y el mismo SHA-256 en ambas herramientas.',
         'successful_primary_gets':34,'insomnia_test_assertions_passed':68,'screenshots':51,
         'body_sha256':body_hash,'python_evidence':py_rel,'insomnia_version':'13.3.0',
         'insomnia_start_local':ins_start,'insomnia_end_local':ins_end,
         'pending':['Identificar y verificar endpoints reales de API','Escalamiento vertical con cliente de prueba','IDOR/BOLA entre dos cuentas propias','Manipulación acotada de IDs propios'],
         'rows':[]}
for r in records:
    summary['rows'].append({'id':r['id'],'endpoint':r['endpoint'],'method':'GET','identity':'Anónimo',
                            'http_python':r['status'],'http_insomnia':ins[r['id']]['statusCode'],
                            'content_type':'text/html','bytes':r['body_bytes'],'sha256':r['sha256'],
                            'result':'Control: HTML común' if r['id'].startswith('CTRL') else 'No concluyente / fallback SPA',
                            'evidence_prefix':r['id'],'insomnia_response_at':local(ins[r['id']]['created'])})
(ROOT/'evidencias/resumen_consolidado.json').write_text(json.dumps(summary,indent=2,ensure_ascii=False)+'\n')
with (ROOT/'evidencias/tabla_resultados_A01.csv').open('w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(summary['rows'][0]));w.writeheader();w.writerows(summary['rows'])

doc=Document();section=doc.sections[0]
section.page_width=Inches(8.27);section.page_height=Inches(11.69)
section.top_margin=Inches(.7);section.bottom_margin=Inches(.65)
section.left_margin=Inches(.7);section.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.05
for style in ['Heading 1','Heading 2','Title']:
    doc.styles[style].font.name='Calibri';doc.styles[style].font.color.rgb=RGBColor.from_string('173D5D')
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Title'].font.size=Pt(24)
header=section.header.paragraphs[0]
header.text='PANIFICADORA NANCY  /  AUDITORÍA DE SISTEMAS  /  A01:2025'
header.style='Caption';header.runs[0].font.size=Pt(8)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Prueba externa · 07/10/2026  |  Página ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title=TITLE
doc.core_properties.subject='Prueba sustantiva de acceso anónimo: 15 rutas, 2 controles y evidencia de Insomnia'
doc.core_properties.author='Ejecución asistida por Codex a solicitud del usuario'

def p(text,style=None):return doc.add_paragraph(text,style)
new_page=False
def h(text,level=1):
    global new_page
    heading=doc.add_heading(text,level)
    if new_page:heading.paragraph_format.page_break_before=True
    new_page=False
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    if widths is None and len(headers)==2:widths=[1.75,5.1]
    if widths:
        t.autofit=False
        for column,width in zip(t.columns,widths):column.width=Inches(width)
    for c,text in zip(t.rows[0].cells,headers):c.text=text
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    for row in t.rows:
        prevent=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(prevent)
        for i,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[i])
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(4)
                par.paragraph_format.space_before=Pt(2)
                for run in par.runs:run.font.size=Pt(9)
    return t

p('INFORME DE PRUEBA SUSTANTIVA','Subtitle')
doc.add_heading(TITLE,0)
p('Fase 1 · Acceso anónimo · Caja negra · DAST','Subtitle')
h('1. Resultado ejecutivo')
p('Resultado general: NO CONCLUYENTE PARA A01. No se confirmó acceso indebido a datos protegidos en las respuestas examinadas. Tampoco existe evidencia suficiente para declarar que el sistema cumple sus controles de acceso.')
p('Las 15 rutas solicitadas respondieron HTTP 200 con Content-Type text/html y 7.541 bytes. El cuerpo es idéntico, byte a byte, al de la página principal y al de una ruta de control diseñada como inexistente. Python Requests e Insomnia 13.3.0 reprodujeron el mismo resultado. Este comportamiento es consistente con el fallback de la SPA.')
p('Se conservaron 34 respuestas HTTP principales satisfactorias, 51 capturas originales de Insomnia y 68 aserciones de instrumentación aprobadas. Estas aserciones comprueban el registro de la respuesta; no son 68 pruebas de autorización ni certifican cumplimiento de OWASP.')
h('2. Datos de la prueba')
table(['Campo','Detalle'],[
    ('Sistema','Panificadora Nancy'),('Destino','https://www.panificadoranancy.com'),
    ('Referencia','OWASP Top 10:2025, A01 — Broken Access Control [R1]'),
    ('Modalidad','Caja negra; DAST acotado; método GET; usuario anónimo'),
    ('Fecha / zona','7 de octubre de 2026 · America/La_Paz (UTC−04:00)'),
    ('Herramientas',f'Insomnia 13.3.0; Python {py["python_version"]}; Requests {py["requests_version"]}; Playwright para operar y capturar la interfaz'),
    ('Responsabilidad','Ejecución técnica asistida por Codex. Responsable académico y revisor: por completar por el usuario.'),
    ('Alcance acordado','Solo acceso anónimo. Las pruebas con cuentas e IDOR se pospusieron por indicación expresa del usuario.')])

new_page=True
h('3. Objetivo, alcance y criterio')
p('Verificar si las rutas indicadas permiten obtener información protegida sin autenticación y diferenciar respuestas de datos de documentos de presentación. La evaluación de funciones y objetos protegidos requiere identificar sus endpoints reales de API y contar con identidades y objetos de prueba controlados.')
p('El alcance comprende 15 rutas proporcionadas por el usuario y 2 controles de comparación. No se consultó el código fuente, configuración del servidor, repositorio o base de datos de la aplicación auditada. Los archivos locales de Insomnia se usaron únicamente para conservar solicitudes, respuestas y resultados de la herramienta.')
table(['Respuesta observada','Interpretación de auditoría'],[
    ('401 / 403','Bloqueo HTTP observado para esa solicitud; no demuestra por sí solo protección de todos los roles, funciones o métodos.'),
    ('404 / 405 / 3xx','Recurso ausente u ocultado, método no permitido o redirección. No permiten un dictamen global. Las redirecciones se mantienen sin seguir.'),
    ('200 con HTML','Inspeccionar el contenido. HTML puede contener datos privados; no se presume automáticamente que sea una SPA.'),
    ('200 con HTML idéntico a ambos controles y contenedor de aplicación vacío','Evidencia de documento común de la SPA. No demuestra acceso al recurso protegido ni valida su API.'),
    ('2xx con JSON u otro contenido','Revisar sensibilidad, política de acceso e identidad. El formato o el éxito HTTP por sí solos no confirman A01.')])
h('4. Procedimiento aplicado')
for text in [
    'Se preparó una colección local separada con 17 solicitudes GET. Sin credenciales, variables con tokens, encabezado Authorization ni Cookie. El envío y almacenamiento de cookies y el seguimiento de redirecciones están desactivados en cada solicitud.',
    'Se ejecutó Python Requests con una sesión nueva por solicitud, trust_env=False, sin credenciales .netrc, con validación TLS y una pausa de 0,8 s entre solicitudes. Se guardaron cuerpos, cabeceras, tiempos y huellas SHA-256.',
    'Se ejecutaron las mismas rutas en Insomnia con Pre-request de limpieza de cabeceras y After-response de clasificación. El timeline confirma la ausencia de Authorization y Cookie; la validación TLS se mantuvo activada.',
    'Se capturaron Preview en modo Source Code, Headers y Tests para cada respuesta. Las imágenes provienen de la interfaz real y no fueron editadas. Se conservaron los resultados de los scripts y la traza HTTP de Insomnia.',
    'Se compararon los bytes de los 34 cuerpos y se revisó el HTML común. Contiene metadatos públicos, JSON-LD promocional, un módulo JavaScript referenciado y un elemento div con id root vacío; no contiene listados privados de pedidos, clientes o inventario.'
]:p(text,'List Number')
p('Incidencia de ejecución: el primer intento de Python falló por resolución DNS dentro del entorno restringido, antes de obtener respuesta HTTP. Se conserva separado y no se incluye en las 34 respuestas válidas. La ejecución con acceso de red autorizado completó el lote.')

new_page=True
h('5. Resultados consolidados por endpoint')
p('Todas las filas corresponden a GET sin autenticación y cuerpos de 7.541 bytes. Criterio esperado: no entregar datos protegidos al usuario anónimo. P/I indica el estado HTTP obtenido en Python e Insomnia. Cada ID vincula tres capturas y las evidencias originales.')
table(['ID','Endpoint','HTTP P/I','Tipo','Resultado'],[
    (r['id'],r['endpoint'],'200/200','text/html','No concluyente / SPA') for r in records if not r['id'].startswith('CTRL')
], [.63,2.45,.68,.83,1.9])
h('5.1. Controles de comparación',2)
table(['ID','Ruta','Obtenido'],[(r['id'],r['endpoint'],'200 HTML / 7.541 bytes') for r in records if r['id'].startswith('CTRL')],[.65,4.1,1.8])
p('Los 34 cuerpos comparten la siguiente huella SHA-256:')
hash_par=p(body_hash);hash_par.runs[0].font.name='Consolas';hash_par.runs[0].font.size=Pt(8)
p('Conclusión técnica: se observó un documento HTML común incluso para el control inexistente. Se infiere un fallback de la SPA a partir del comportamiento externo; no se afirma haber inspeccionado la configuración del servidor.')

new_page=True
h('6. Hallazgos, límites y pruebas pendientes')
p('OBS-A01-01 — Documento HTML común en rutas de navegación. Tipo: observación técnica, sin vulnerabilidad A01 confirmada. Severidad/CVSS: no asignados. Evidencia: matriz de resultados, comparación de cuerpos y capturas asociadas.')
p('El HTTP 200 observado no prueba que un usuario anónimo haya obtenido un pedido o información administrativa. No corresponde clasificar estas respuestas como “No cumple”. La ausencia de datos privados en el documento común tampoco acredita que la API aplique controles de acceso correctos; por ello no corresponde clasificarlas como “Cumple”.')
h('6.1. Cobertura no evaluada',2)
table(['Prueba pendiente','Estado y condición necesaria'],[
    ('Acceso anónimo a APIs reales de datos','Pendiente: identificar URLs y métodos a partir del tráfico normal del cliente o de documentación autorizada; las rutas de navegación no identifican por sí solas las APIs.'),
    ('Escalamiento vertical / BFLA','No ejecutado por alcance acordado. Requiere una cuenta cliente válida y una función administrativa identificada. Primero confirmar que la sesión sea válida mediante un control positivo.'),
    ('IDOR / BOLA entre A y B','No ejecutado por alcance acordado. Requiere dos cuentas propias y un objeto conocido de cada una; probar la lectura cruzada y contrastar con la lectura legítima.'),
    ('Manipulación de identificadores','No ejecutada. Usar identificadores de objetos de prueba autorizados; no se realizó un barrido de IDs ajenos.'),
    ('Otros métodos y escenarios','No evaluados: POST, PUT, PATCH, DELETE, sesiones autenticadas, cambios de roles, escritura, eliminación, endpoints no incluidos ni configuración interna.')])
h('6.2. Recomendaciones',2)
for text in [
    'Continuar en el mismo informe con endpoints de API observados y una matriz de permisos esperados por identidad, función y objeto.',
    'Comprobar la autorización en el servidor para cada función y cada objeto protegido, además de cualquier restricción visual del frontend. Este es un criterio de evaluación recomendado por OWASP [R1], no una deficiencia comprobada en este sistema.',
    'Cuando se amplíe el alcance, comparar la respuesta legítima del propietario con la respuesta de otra cuenta de prueba y con la del usuario anónimo. Conservar códigos, contenido pertinente y contexto sin incluir tokens en el informe.',
    'Mantener el estado global “No concluyente” hasta reunir evidencia suficiente. Un acceso indebido confirmado justificaría “No cumple” para el control afectado; completar esta muestra sin hallazgos no equivale a certificar todo el sistema.'
]:p(text,'List Bullet')

new_page=True
h('7. Evidencias y trazabilidad')
p(f'Python: {py_start} a {py_end}. Insomnia (fecha de registro de respuestas): {ins_start} a {ins_end}. Las horas de cada captura se encuentran en el índice, con desplazamiento UTC−04:00.')
table(['Archivo o carpeta','Contenido'],[
    ('capturas/','51 PNG originales: 3 por cada una de las 15 rutas y los 2 controles. Ejemplo: A01-02_preview.png, A01-02_headers.png y A01-02_tests.png.'),
    ('capturas/indice_capturas.json','Relación URL, ID de solicitud/respuesta de Insomnia, vista, fecha de captura y hashes de respuesta e imagen.'),
    ('evidencias/insomnia/','Por cada ID: cuerpo .body, cabeceras/resultados .response.json, solicitud .request.json, traza .timeline y textos de accesibilidad .ui.txt.'),
    ('Lote Python válido', py_rel+'/ — Resultados TXT, CSV y JSON; cuerpo y metadatos por solicitud.'),
    ('Resumen y tabla CSV','evidencias/resumen_consolidado.json y evidencias/tabla_resultados_A01.csv: comparación entre herramientas.'),
    ('Colección de Insomnia','coleccion_insomnia/A01_Nancy_anonimo.insomnia.json. Incluye Pre-request y After-response, sin credenciales.'),
    ('scripts/','Enumeración acotada, scripts de Insomnia, automatización de capturas y generación del informe.'),
    ('MANIFIESTO_SHA256.txt','Huellas de los archivos entregables para detectar cambios posteriores. No constituye certificación de fecha por un tercero.')])
h('7.1. Lectura correcta de las capturas',2)
p('Preview muestra el HTML recibido en modo Source Code, sin confundir la representación visual con datos de API. Headers muestra Content-Type y las cabeceras del servidor. Tests muestra 4/4 aserciones de registro aprobadas; la propia etiqueta aclara “no certifica A01”.')
p('La pestaña Auth muestra “Inherit from parent”; el perfil Scratch Pad no tiene credenciales de padre ni entorno con secretos. La ausencia efectiva de autenticación está comprobada en las cabeceras salientes del timeline y en la configuración de cada solicitud; no se deduce únicamente del icono de la interfaz.')
p('Las imágenes seleccionadas de las páginas siguientes son una muestra. La cobertura completa está en las 51 capturas externas. Para inspeccionar el texto pequeño, abrir el PNG original o ampliar el PDF. El informe previo hallado en la carpeta se conservó sin cambios y no se usó como evidencia de Panificadora Nancy.')
h('8. Referencias y cierre')
p('[R1] OWASP Top 10:2025 — A01 Broken Access Control. https://top10.owasp.org/2025/A01_2025-Broken_Access_Control/ (consulta: 07/10/2026).')
p('[R2] Kong — Scripts in Insomnia. https://developer.konghq.com/insomnia/scripts/ (consulta: 07/10/2026). Se emplearon sus interfaces de respuesta y pruebas para registrar la evidencia.')
p('Dictamen de esta fase: NO CONCLUYENTE PARA A01. Fallback HTML observado; ningún acceso indebido confirmado en las respuestas estudiadas; control de acceso de APIs y sesiones pendientes de evaluación.')

chosen=[
    ('A01-02_preview.png','E1 · /mis-pedidos — Preview del HTML','La ruta devuelve un documento HTML con metadatos de la aplicación. El contenido completo se conserva como A01-02.body.'),
    ('A01-02_headers.png','E2 · /mis-pedidos — Cabeceras','HTTP 200 y Content-Type text/html; el tamaño íntegro observado es de 7.541 bytes.'),
    ('A01-02_tests.png','E3 · /mis-pedidos — Tests After-response','Los cuatro resultados PASS corresponden a instrumentación. La clasificación indica que el HTML no es concluyente para A01.'),
    ('A01-04_tests.png','E4 · /admin/usuarios — Tests After-response','Acceso anónimo a una ruta de navegación administrativa. El cuerpo coincide con el documento de la SPA, sin demostrar acceso a una API administrativa.'),
    ('A01-12_headers.png','E5 · /inventario/movimientos-caja — Cabeceras','La ruta de inventario muestra el mismo comportamiento HTTP/HTML. No se obtuvo un listado de movimientos de caja en el cuerpo observado.'),
    ('CTRL-02_preview.png','E6 · Ruta de control inexistente — Preview','También recibe el documento común. La igualdad de SHA-256 con raíz y rutas evaluadas sustenta la inferencia de fallback.')]
sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
sec.top_margin=Inches(.5);sec.bottom_margin=Inches(.5);sec.left_margin=Inches(.5);sec.right_margin=Inches(.5)
for i,(filename,title,caption) in enumerate(chosen):
    if i:new_page=True
    h(title)
    p(caption)
    doc.add_picture(str(ROOT/'capturas'/filename),width=Inches(10.35))
    doc.paragraphs[-1].paragraph_format.line_spacing=1.0
    doc.paragraphs[-1].paragraph_format.space_after=Pt(0)
    meta=next(c for c in captures if c['file']==filename)
    p('Original: capturas/'+filename+' · Capturada: '+meta['captured_at'],'Caption')
doc.save(out/'Informe_A01_Controles_Acceso_Nancy.docx')

md=['# '+TITLE,'','**Resultado: No concluyente para A01.**','',
    'Fase ejecutada: GET anónimo, caja negra. Fecha: 07/10/2026, America/La_Paz.',
    '','Las 15 rutas y los 2 controles devuelven HTTP 200, text/html, 7.541 bytes y el mismo cuerpo en Insomnia y Python. La respuesta es consistente con el fallback de la SPA. No se confirmó acceso indebido ni se certificó la autorización de la API.',
    '',f'SHA-256 común: `{body_hash}`.','',
    '34 respuestas principales válidas; 51 capturas originales; 68 aserciones de instrumentación aprobadas (no son pruebas de cumplimiento A01).','',
    '| ID | Endpoint | HTTP Python / Insomnia | Resultado |','|---|---|---|---|']
for r in summary['rows']:md.append(f'| {r["id"]} | `{r["endpoint"]}` | 200 / 200 | {r["result"]} |')
md+=['','Pendientes por el alcance acordado: APIs reales de datos, cliente frente a funciones administrativas, IDOR/BOLA con dos cuentas propias y manipulación de IDs controlados.',
     '','El detalle completo, limitaciones y muestra de capturas se encuentran en el DOCX y PDF de este mismo informe.',
     '',f'Evidencia Python: `../{py_rel}/`. Evidencia Insomnia: `../evidencias/insomnia/`. Capturas: `../capturas/`.',
     '','Criterio de autorización: [OWASP A01:2025](https://top10.owasp.org/2025/A01_2025-Broken_Access_Control/). Scripts de respuesta: [Kong Insomnia](https://developer.konghq.com/insomnia/scripts/).']
(out/'Informe_A01_Controles_Acceso_Nancy.md').write_text('\n'.join(md)+'\n')
print(out/'Informe_A01_Controles_Acceso_Nancy.docx')
print(json.dumps({k:v for k,v in summary.items() if k not in ['rows','pending']},ensure_ascii=False,indent=2))
